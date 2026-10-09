"""Team Tracker API.

Run:  uvicorn app.main:app --reload   (from the backend/ directory)

Mutations return the full, freshly annotated snapshot so the client can simply
replace its state. Writes that the UI labels "Written to Jira" go through here; the
store is the integration seam where a real Jira client would push changes.
"""

import csv
import io
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from .dates import days_between, dm, js_weekday
from .domain import LEAVE_TYPES, STEPS, answer, find_item, validate_leave_rows
from .seed import DEFAULT_RULE, EPIC_CATALOG, TODAY
from .store import Store

BASE = Path(__file__).resolve().parent.parent
DATA_PATH = Path(os.environ.get("TRACKER_DATA", BASE / "data" / "store.json"))
DIST = BASE.parent / "frontend" / "dist"

store = Store(DATA_PATH, os.environ.get("TRACKER_TODAY", TODAY))
app = FastAPI(title="Team Tracker API")

ISO_DATE = r"^\d{4}-\d{2}-\d{2}$"
Status = Literal["N", "W", "R", "C"]


def _member(s: dict, mid: str):
    if mid not in s["members"]:
        raise HTTPException(400, f"Unknown member {mid}")


# ---------------------------------------------------------------- state & sync

@app.get("/api/state")
def get_state():
    return store.snapshot()


@app.post("/api/sync")
def sync():
    store.touch_sync()
    return store.snapshot()


@app.post("/api/reset")
def reset():
    store.reset()
    return store.snapshot()


# ---------------------------------------------------------------- tickets

class TicketPatch(BaseModel):
    status: Status | None = None
    assignee: str | None = None
    ps: str | None = Field(None, pattern=ISO_DATE)
    pe: str | None = Field(None, pattern=ISO_DATE)


@app.get("/api/tickets/{key}")
def get_ticket(key: str):
    c = find_item(store.snapshot(), key.strip().upper())
    if not c:
        raise HTTPException(404, f"Ticket {key or '(empty)'} was not found")
    return c


@app.patch("/api/tickets/{key}")
def patch_ticket(key: str, body: TicketPatch):
    with store.mutate() as s:
        c = find_item(s, key)
        if not c:
            raise HTTPException(404, f"Ticket {key} was not found")
        changes = body.model_dump(exclude_none=True)
        if "assignee" in changes:
            _member(s, changes["assignee"])
            c["held"] = 0
        c.update(changes)
    return store.snapshot()


class PlanChange(BaseModel):
    key: str
    ps: str = Field(pattern=ISO_DATE)
    pe: str = Field(pattern=ISO_DATE)
    who: str


class PlanPush(BaseModel):
    changes: list[PlanChange]


@app.post("/api/plan/push")
def push_plan(body: PlanPush):
    with store.mutate() as s:
        for ch in body.changes:
            c = find_item(s, ch.key)
            if not c:
                raise HTTPException(404, f"Ticket {ch.key} was not found")
            _member(s, ch.who)
            c["ps"], c["pe"] = ch.ps, ch.pe
            if c["worker"] != ch.who:
                c["worker"] = c["assignee"] = ch.who
                c["held"] = 0
    return store.snapshot()


class NewStep(BaseModel):
    ps: str = Field(pattern=ISO_DATE)
    pe: str = Field(pattern=ISO_DATE)
    est: float = Field(ge=0)
    who: str


class NewFunction(BaseModel):
    parent: str
    phase: str
    fid: str = Field(pattern=r"^[A-Z]\d+$")
    name: str = Field(min_length=1)
    ba: str
    rows: list[NewStep] = Field(min_length=5, max_length=5)


@app.post("/api/functions")
def create_function(body: NewFunction):
    with store.mutate() as s:
        if any(f["id"] == body.fid for f in s["functions"]):
            raise HTTPException(409, f"Function {body.fid} already exists")
        _member(s, body.ba)
        cells = []
        for i, r in enumerate(body.rows):
            _member(s, r.who)
            st = STEPS[i]
            cells.append({
                "key": f"PAY-{s['nextKey']}", "fid": body.fid, "fname": body.name.strip(), "ba": body.ba, "si": i,
                "sk": st["k"], "stepLabel": st["label"], "isFix": bool(st.get("fix")), "isRev": bool(st.get("rev")),
                "phase": body.phase, "status": "N", "logged": 0, "est": r.est, "assignee": r.who, "worker": r.who,
                "ps": r.ps, "pe": max(r.ps, r.pe), "rounds": 0, "held": 0, "prj": "PAY", "worklogs": [],
            })
            s["nextKey"] += 1
        s["functions"].append({"id": body.fid, "name": body.name.strip(), "ba": body.ba, "cells": cells})
    return store.snapshot()


# ---------------------------------------------------------------- Q&A

class Question(BaseModel):
    q: str = Field(min_length=1)


@app.post("/api/chat")
def chat(body: Question):
    return answer(store.snapshot(), body.q)


# ---------------------------------------------------------------- leave

class LeaveIn(BaseModel):
    editId: int | None = None
    who: str
    type: Literal["full", "half", "fullTent", "halfTent"]
    from_: str = Field(alias="from", pattern=ISO_DATE)
    to: str = Field(pattern=ISO_DATE)
    note: str = ""


def _visible(s: dict, d: str) -> bool:
    return js_weekday(d) in s["config"]["days"] or d in s["extraDays"]


@app.post("/api/leaves")
def save_leave(body: LeaveIn):
    if body.to < body.from_:
        raise HTTPException(400, "To date is before From date")
    with store.mutate() as s:
        _member(s, body.who)
        if body.editId is not None:
            s["leaves"] = [l for l in s["leaves"] if l["id"] != body.editId]
        n = 0
        for d in days_between(body.from_, body.to):
            if n >= 60:
                break
            if not _visible(s, d):
                continue
            ex = next((l for l in s["leaves"] if l["who"] == body.who and l["date"] == d), None)
            if ex:
                ex.update(type=body.type, note=body.note)
            else:
                s["leaves"].append({"id": s["nextLeaveId"], "who": body.who, "date": d, "type": body.type, "note": body.note})
                s["nextLeaveId"] += 1
            n += 1
    return store.snapshot()


@app.delete("/api/leaves/{lid}")
def delete_leave(lid: int):
    with store.mutate() as s:
        s["leaves"] = [l for l in s["leaves"] if l["id"] != lid]
    return store.snapshot()


IMPORT_SAMPLE = [
    ["An", "2026-11-05", "2026-11-06", "All day", "Family trip"], ["Bình", "2026-11-10", "2026-11-10", "All day", ""],
    ["Chi", "2026-11-12", "2026-11-11", "All day", ""], ["Duc", "2026-11-13", "2026-11-13", "Half day (tentative)", ""],
    ["Dũng", "2026-11-16", "2026-11-16", "Half day", "Morning off"], ["Minh", "2026-11-18", "2026-11-20", "All day", ""],
    ["Tanaka", "2026-11-24", "2026-11-24", "Half day (tentative)", ""], ["An", "2026-10-16", "2026-10-16", "All day", ""],
    ["Bình", "2026/11/25", "2026/11/25", "All day", ""], ["Chi", "2026-11-26", "2026-11-27", "Remote", ""],
]


def _cell(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return "" if v is None else str(v).strip()


def _read_rows(name: str, raw: bytes) -> list[list[str]]:
    if name.lower().endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        rows = [r for r in csv.reader(io.StringIO(text)) if any(x.strip() for x in r)]
        return [[x.strip() for x in r] for r in rows[1:]]
    from openpyxl import load_workbook

    try:
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
    except Exception as e:  # noqa: BLE001 - surface any parse failure as a 400
        raise HTTPException(400, f"Could not read {name}: {e}") from e
    rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    return [r for r in rows[1:] if any(r)]


def _preview(name: str, rows: list[list[str]]) -> dict:
    return {"file": name, "rows": validate_leave_rows(store.snapshot(), rows)}


@app.post("/api/leaves/import/preview")
async def import_preview(file: UploadFile = File(...)):
    name = file.filename or "upload"
    if not re.search(r"\.(csv|xlsx)$", name, re.I):
        raise HTTPException(400, "Choose a .xlsx or .csv file")
    return _preview(name, _read_rows(name, await file.read()))


@app.get("/api/leaves/import/sample")
def import_sample():
    return _preview("leave-november.xlsx", IMPORT_SAMPLE)


class ImportConfirm(BaseModel):
    file: str
    rows: list[list[str]]


@app.post("/api/leaves/import")
def import_confirm(body: ImportConfirm):
    checked = validate_leave_rows(store.snapshot(), body.rows)
    ok = [r for r in checked if not r["err"]]
    n = 0
    with store.mutate() as s:
        for r in ok:
            for d in days_between(r["from"], r["to"]):
                if js_weekday(d) not in s["config"]["days"]:
                    continue
                s["leaves"].append({"id": s["nextLeaveId"], "who": r["id"], "date": d, "type": r["tk"], "note": r["note"]})
                s["nextLeaveId"] += 1
                n += 1
    return {"imported": n, "month": ok[0]["from"][:7] if ok else None, "snapshot": store.snapshot()}


@app.get("/api/leaves/template")
def leave_template():
    body = ("﻿Member,From (YYYY-MM-DD),To (YYYY-MM-DD),Type,Note\nAn,2026-11-05,2026-11-06,All day,Family trip\n"
            "Bình,2026-11-10,2026-11-10,Half day,\nDũng,2026-11-16,2026-11-16,All day (tentative),\n"
            "Tanaka,2026-11-24,2026-11-24,Half day (tentative),\n")
    return Response(body, media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="leave-import-template.csv"'})


@app.get("/api/leaves/export")
def leave_export(month: str = Query(pattern=r"^\d{4}-\d{2}$")):
    s = store.snapshot()
    esc = lambda x: str(x).replace("&", "&amp;").replace("<", "&lt;")  # noqa: E731
    leaves = sorted((l for l in s["leaves"] if l["date"].startswith(month)), key=lambda l: l["date"])
    hols = [d for d in sorted(s["holidays"]) if d.startswith(month)]
    rows = "".join(f"<tr><td>{esc(s['members'][l['who']]['full'])}</td><td>{l['date']}</td>"
                   f"<td>{LEAVE_TYPES[l['type']]['label']}</td><td>{esc(l['note'])}</td></tr>" for l in leaves)
    rows += "".join(f"<tr><td>{s['holidays'][d][0]}</td><td>{d}</td><td>Holiday</td><td>{esc(s['holidays'][d][1])}</td></tr>" for d in hols)
    html = ('<html><head><meta charset="utf-8"></head><body><table border="1"><tr><th>Member</th><th>Date</th>'
            f"<th>Type</th><th>Note</th></tr>{rows}</table></body></html>")
    return Response(html, media_type="application/vnd.ms-excel",
                    headers={"Content-Disposition": f'attachment; filename="team-calendar-{month}.xls"'})


# ---------------------------------------------------------------- calendar settings

class HolidayIn(BaseModel):
    date: str = Field(pattern=ISO_DATE)
    name: str = Field(min_length=1)
    c: Literal["VN", "JP", "VN JP"]


@app.post("/api/holidays")
def add_holiday(body: HolidayIn):
    with store.mutate() as s:
        s["holidays"][body.date] = [body.c, body.name.strip()]
    return store.snapshot()


@app.delete("/api/holidays/{d}")
def delete_holiday(d: str):
    with store.mutate() as s:
        s["holidays"].pop(d, None)
    return store.snapshot()


class ExtraDayIn(BaseModel):
    date: str = Field(pattern=ISO_DATE)
    note: str = ""


@app.post("/api/extra-days")
def add_extra_day(body: ExtraDayIn):
    with store.mutate() as s:
        if js_weekday(body.date) in s["config"]["days"]:
            raise HTTPException(400, f"{dm(body.date)} is already a working day")
        if body.date in s["extraDays"]:
            raise HTTPException(400, f"{dm(body.date)} is already in the list")
        s["extraDays"][body.date] = body.note.strip() or "Extra working day"
    return store.snapshot()


@app.delete("/api/extra-days/{d}")
def delete_extra_day(d: str):
    with store.mutate() as s:
        s["extraDays"].pop(d, None)
    return store.snapshot()


class ConfigPatch(BaseModel):
    hours: float | None = Field(None, gt=0, le=24)
    tsTol: float | None = Field(None, ge=0)
    soon: float | None = Field(None, ge=0)
    effWarn: float | None = Field(None, ge=0)
    held: float | None = Field(None, ge=0)
    rounds: float | None = Field(None, ge=0)
    cap: float | None = Field(None, ge=0)
    days: list[int] | None = None


@app.patch("/api/config")
def patch_config(body: ConfigPatch):
    changes = body.model_dump(exclude_none=True)
    if "days" in changes and not all(0 <= d <= 6 for d in changes["days"]):
        raise HTTPException(400, "Working days must be 0–6")
    with store.mutate() as s:
        s["config"].update(changes)
    return store.snapshot()


class MemberPatch(BaseModel):
    role: Literal["BA", "BAL", "JM"] | None = None
    perm: Literal["admin", "user"] | None = None


@app.patch("/api/members/{mid}")
def patch_member(mid: str, body: MemberPatch):
    with store.mutate() as s:
        _member(s, mid)
        s["members"][mid].update(body.model_dump(exclude_none=True))
    return store.snapshot()


# ---------------------------------------------------------------- Jira settings

class JiraIn(BaseModel):
    url: str
    email: str
    token: str | None = None  # omitted = keep the stored token
    interval: Literal["5", "15", "60", "manual"]


def _jira_ok(url: str, token: str) -> bool:
    # Simulated check; a live client would call GET {url}/rest/api/3/myself with these credentials.
    return bool(re.search(r"atlassian\.net", url)) and len(token) > 8


@app.put("/api/settings/jira")
def save_jira(body: JiraIn):
    with store.mutate() as s:
        j = s["settings"]["jira"]
        j.update(url=body.url.strip(), email=body.email.strip(), interval=body.interval)
        if body.token:
            j["token"] = body.token
        j["status"] = "ok" if _jira_ok(j["url"], j["token"]) else "fail"
    return store.snapshot()


class JiraTest(BaseModel):
    url: str
    token: str | None = None


@app.post("/api/settings/jira/test")
def test_jira(body: JiraTest):
    token = body.token or store.state["settings"]["jira"]["token"]
    return {"status": "ok" if _jira_ok(body.url, token) else "fail"}


class EpicIn(BaseModel):
    key: str


@app.post("/api/settings/epics")
def add_epic(body: EpicIn):
    k = body.key.strip().upper()
    if not k:
        raise HTTPException(400, "Enter an Epic key such as PAY-1100")
    with store.mutate() as s:
        epics = s["settings"]["jira"]["epics"]
        if k in epics:
            raise HTTPException(409, f"{k} is already tracked")
        if k not in EPIC_CATALOG:
            msg = f"{k} was not found or is not an Epic" if re.match(r"^[A-Z]+-\d+$", k) else "Enter an Epic key such as PAY-1100"
            raise HTTPException(404, msg)
        epics.append(k)
    return store.snapshot()


@app.delete("/api/settings/epics/{key}")
def delete_epic(key: str):
    with store.mutate() as s:
        j = s["settings"]["jira"]
        j["epics"] = [x for x in j["epics"] if x != key]
    return store.snapshot()


class SettingsPatch(BaseModel):
    map: dict[str, str] | None = None
    rule: str | None = None
    alias: dict[str, str] | None = None


@app.patch("/api/settings")
def patch_settings(body: SettingsPatch):
    with store.mutate() as s:
        st = s["settings"]
        if body.map is not None:
            st["map"].update({k: v for k, v in body.map.items() if k in ("ps", "pe", "eff")})
        if body.rule is not None:
            st["rule"] = body.rule or DEFAULT_RULE
        if body.alias is not None:
            st["alias"].update({k: v for k, v in body.alias.items() if k in st["alias"]})
    return store.snapshot()


# ---------------------------------------------------------------- static frontend (production build)

if DIST.exists():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")

    @app.get("/{path:path}", include_in_schema=False)
    def spa(path: str):
        f = DIST / path
        return FileResponse(f if path and f.is_file() else DIST / "index.html")

