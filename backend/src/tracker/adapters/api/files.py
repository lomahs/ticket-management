"""Spreadsheet I/O for leave import/export: reading CSV/XLSX rows and writing the Excel-readable export."""

import csv
import io
from datetime import date, datetime
from html import escape

from tracker.application.errors import Invalid

TEMPLATE_CSV = ("﻿Member,From (YYYY-MM-DD),To (YYYY-MM-DD),Type,Note\n"
                "An,2026-11-05,2026-11-06,All day,Family trip\nBình,2026-11-10,2026-11-10,Half day,\n"
                "Dũng,2026-11-16,2026-11-16,All day (tentative),\nTanaka,2026-11-24,2026-11-24,Half day (tentative),\n")

SAMPLE_FILE = "leave-november.xlsx"
SAMPLE_ROWS = [
    ["An", "2026-11-05", "2026-11-06", "All day", "Family trip"], ["Bình", "2026-11-10", "2026-11-10", "All day", ""],
    ["Chi", "2026-11-12", "2026-11-11", "All day", ""], ["Duc", "2026-11-13", "2026-11-13", "Half day (tentative)", ""],
    ["Dũng", "2026-11-16", "2026-11-16", "Half day", "Morning off"], ["Minh", "2026-11-18", "2026-11-20", "All day", ""],
    ["Tanaka", "2026-11-24", "2026-11-24", "Half day (tentative)", ""], ["An", "2026-10-16", "2026-10-16", "All day", ""],
    ["Bình", "2026/11/25", "2026/11/25", "All day", ""], ["Chi", "2026-11-26", "2026-11-27", "Remote", ""],
]


def _cell(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return "" if v is None else str(v).strip()


def read_rows(name: str, raw: bytes) -> list[list[str]]:
    """Data rows (header skipped) of the first sheet of an .xlsx, or of a .csv."""
    lower = name.lower()
    if lower.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        rows = [r for r in csv.reader(io.StringIO(text)) if any(x.strip() for x in r)]
        return [[x.strip() for x in r] for r in rows[1:]]
    if lower.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except Exception as e:  # noqa: BLE001 - any parse failure is a bad upload
            raise Invalid(f"Could not read {name}: {e}", "file") from e
        rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
        return [r for r in rows[1:] if any(r)]
    raise Invalid("Choose a .xlsx or .csv file", "file")


def export_xls(rows: list[tuple[str, str, str, str]]) -> str:
    """An HTML table Excel opens as a sheet (no extra dependency, keeps Vietnamese/Japanese text intact)."""
    body = "".join(f"<tr><td>{escape(a)}</td><td>{b}</td><td>{escape(c)}</td><td>{escape(d)}</td></tr>" for a, b, c, d in rows)
    return ('<html><head><meta charset="utf-8"></head><body><table border="1"><tr><th>Member</th><th>Date</th><th>Type</th>'
            f"<th>Note</th></tr>{body}</table></body></html>")
